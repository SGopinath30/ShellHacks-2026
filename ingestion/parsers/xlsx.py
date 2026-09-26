from io import BytesIO

from openpyxl import load_workbook


def parse_xlsx(data):
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
    try:
        return {"format": "xlsx", "sheets": [
            {"name": sheet.title, "rows": [
                {"row_number": row[0].row if row and hasattr(row[0], "row") else number,
                 "values": [cell.value for cell in row],
                 "cell_types": [cell.data_type for cell in row],
                 "number_formats": [cell.number_format for cell in row]}
                for number, row in enumerate(sheet.iter_rows(), 1)
                if any(cell.value is not None for cell in row)]}
            for sheet in workbook.worksheets]}
    finally:
        workbook.close()
