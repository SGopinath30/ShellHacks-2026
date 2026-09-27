"""Deterministic starter-workbook extraction with cell-level evidence."""

from __future__ import annotations

import hashlib
import re
from collections.abc import Iterable, Mapping
from pathlib import Path
from typing import Any

from project_intelligence.contracts import (
    AssociationState,
    CandidateProject,
    CanonicalProjectStatus,
    EvidenceLocator,
    EvidenceValue,
    ExtractionMethod,
    FieldEvidence,
    ProjectType,
    SourceAccess,
    SourceArtifact,
    SourceType,
    VerificationState,
    VersionState,
)
from project_intelligence.normalization.projects import (
    normalize_project_name,
    normalize_utility,
)
from project_intelligence.normalization.schedules import normalize_schedule
from project_intelligence.normalization.status import normalize_status
from project_intelligence.normalization.voltage import normalize_voltage_kv
from project_intelligence.source_safety import RestrictedSourceError


HEADER_ALIASES = {
    "utility": {"utility", "utility name", "company"},
    "name": {"project", "project name", "project_name", "name"},
    "description": {"description", "project description", "scope", "project scope"},
    "external_project_id": {
        "project id",
        "project_id",
        "external project id",
        "utility project id",
    },
    "voltage_kv": {"voltage", "voltage kv", "voltage_kv", "kv"},
    "status": {"status", "project status"},
    "project_type": {"project type", "project_type", "type"},
    "in_service_date": {
        "in service date",
        "in-service date",
        "in_service_date",
        "commercial operation date",
        "cod",
        "projected in service",
    },
    "need_date": {"need date", "need_date"},
    "endpoint_a": {"endpoint a", "endpoint_a", "from", "terminal a"},
    "endpoint_b": {"endpoint b", "endpoint_b", "to", "terminal b"},
    "state": {"state"},
    "county": {"county"},
}


def parse_starter_workbook(source: SourceArtifact) -> list[CandidateProject]:
    if source.source_type is not SourceType.STARTER_WORKBOOK:
        raise ValueError("starter workbook parser requires STARTER_WORKBOOK source type")
    if source.source_access is SourceAccess.CEII:
        raise RestrictedSourceError("CEII starter workbooks are rejected")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("XLSX parsing requires openpyxl") from exc

    workbook_path = Path(source.path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    candidates: list[CandidateProject] = []
    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows())
        if not rows:
            continue
        headers = [str(cell.value or "").strip() for cell in rows[0]]
        column_lookup = _column_lookup(headers)
        for row_number, cells in enumerate(rows[1:], start=2):
            row = {
                field: cells[column_index].value
                for field, column_index in column_lookup.items()
                if column_index < len(cells)
            }
            columns = {
                field: cells[column_index].column_letter
                for field, column_index in column_lookup.items()
                if column_index < len(cells)
            }
            candidate = _parse_row(
                row,
                source=source,
                sheet=worksheet.title,
                row_number=row_number,
                columns=columns,
                labels={
                    field: headers[column_index]
                    for field, column_index in column_lookup.items()
                },
            )
            if candidate:
                candidates.append(candidate)
    return candidates


def parse_starter_rows(
    rows: Iterable[Mapping[str, Any]],
    *,
    source: SourceArtifact,
    sheet: str = "Projects",
    first_data_row: int = 2,
) -> list[CandidateProject]:
    if source.source_access is SourceAccess.CEII:
        raise RestrictedSourceError("CEII starter data is rejected")
    candidates = []
    for offset, supplied_row in enumerate(rows):
        normalized_row = {
            _canonical_header(header): value for header, value in supplied_row.items()
        }
        candidate = _parse_row(
            normalized_row,
            source=source,
            sheet=sheet,
            row_number=first_data_row + offset,
            columns={field: field for field in normalized_row},
            labels={field: field for field in normalized_row},
        )
        if candidate:
            candidates.append(candidate)
    return candidates


def _parse_row(
    row: Mapping[str, Any],
    *,
    source: SourceArtifact,
    sheet: str,
    row_number: int,
    columns: Mapping[str, str],
    labels: Mapping[str, str],
) -> CandidateProject | None:
    raw_name = row.get("name")
    if raw_name is None or not str(raw_name).strip():
        return None
    evidence: list[FieldEvidence] = []

    def bind(
        field_path: str,
        raw_value: object,
        normalized_value: object,
        *,
        state: VerificationState = VerificationState.VERIFIED_RULE,
    ) -> str:
        evidence_id = _evidence_id(source.source_version_id, sheet, row_number, field_path)
        evidence.append(
            FieldEvidence(
                evidence_id=evidence_id,
                field_path=field_path,
                source_version_id=source.source_version_id,
                locator=EvidenceLocator(
                    sheet=sheet,
                    row=row_number,
                    column=columns.get(field_path.split(".")[0]),
                    element_id=f"{sheet}!{row_number}",
                ),
                original_text=str(raw_value),
                normalized_value=_serializable_value(normalized_value),
                extraction_method=ExtractionMethod.TABLE_PARSER,
                validation_state=state,
                association_state=AssociationState.ASSOCIATION_VERIFIED,
            )
        )
        return evidence_id

    name = normalize_project_name(raw_name)
    name_evidence = bind("name", raw_name, name)
    raw_utility = row.get("utility") or source.utility
    utility = normalize_utility(raw_utility)
    utility_evidence = bind("utility_id", raw_utility, utility)
    candidate_seed = f"{source.source_version_id}|{sheet}|{row_number}|{utility}|{name}"
    candidate_id = f"CP-{hashlib.sha256(candidate_seed.encode()).hexdigest()[:12]}"

    raw_voltage = row.get("voltage_kv")
    voltage = normalize_voltage_kv(raw_voltage)
    voltage_value = None
    if raw_voltage not in (None, ""):
        voltage_state = (
            VerificationState.VERIFIED_RULE
            if voltage is not None
            else VerificationState.UNRESOLVED
        )
        voltage_evidence = bind(
            "voltage_kv", raw_voltage, voltage, state=voltage_state
        )
        voltage_value = EvidenceValue[float](
            value=voltage,
            raw=str(raw_voltage),
            state=voltage_state,
            evidence_ids=[voltage_evidence],
        )

    description_value = _text_value("description", row.get("description"), bind)

    raw_status = row.get("status")
    status_value = None
    if raw_status not in (None, ""):
        status = normalize_status(raw_status)
        status_state = (
            VerificationState.VERIFIED_RULE
            if status is not CanonicalProjectStatus.UNKNOWN
            else VerificationState.UNRESOLVED
        )
        status_evidence = bind("status", raw_status, status.value, state=status_state)
        status_value = EvidenceValue[CanonicalProjectStatus](
            value=(status if status is not CanonicalProjectStatus.UNKNOWN else None),
            raw=str(raw_status),
            state=status_state,
            evidence_ids=[status_evidence],
        )

    project_type_value = _project_type_value(row.get("project_type"), bind)
    schedule = None
    for schedule_field in ("in_service_date", "need_date"):
        raw_schedule = row.get(schedule_field)
        if raw_schedule not in (None, ""):
            schedule_evidence = bind(schedule_field, raw_schedule, raw_schedule)
            schedule = normalize_schedule(
                raw_schedule,
                label=labels.get(schedule_field, schedule_field),
                evidence_ids=[schedule_evidence],
            )
            break

    endpoints = []
    for endpoint_field in ("endpoint_a", "endpoint_b"):
        raw_endpoint = row.get(endpoint_field)
        if raw_endpoint not in (None, ""):
            endpoint = normalize_project_name(raw_endpoint)
            endpoint_evidence = bind(endpoint_field, raw_endpoint, endpoint)
            endpoints.append(
                EvidenceValue[str](
                    value=endpoint,
                    raw=str(raw_endpoint),
                    state=VerificationState.VERIFIED_RULE,
                    evidence_ids=[endpoint_evidence],
                )
            )

    state_value = _text_value("state", row.get("state"), bind)
    county_value = _text_value("county", row.get("county"), bind)
    external_project_id = _optional_text(row.get("external_project_id"))
    external_project_id_evidence = []
    if external_project_id:
        external_project_id_evidence = [
            bind(
                "external_project_id",
                row.get("external_project_id"),
                external_project_id,
            )
        ]

    return CandidateProject(
        candidate_project_id=candidate_id,
        source_version_id=source.source_version_id,
        source_access=source.source_access,
        utility_id=utility,
        utility_id_evidence_ids=[utility_evidence],
        external_project_id=external_project_id,
        external_project_id_evidence_ids=external_project_id_evidence,
        name=EvidenceValue[str](
            value=name,
            raw=str(raw_name),
            state=VerificationState.VERIFIED_RULE,
            evidence_ids=[name_evidence],
        ),
        description=description_value,
        voltage_kv=voltage_value,
        status=status_value,
        project_type=project_type_value,
        schedule=schedule,
        endpoints=endpoints,
        state=state_value,
        county=county_value,
        field_evidence=evidence,
        version_state=VersionState.STARTER_DATA,
    )


def _project_type_value(raw_value: object, bind) -> EvidenceValue[ProjectType] | None:
    if raw_value in (None, ""):
        return None
    normalized = re.sub(r"[^a-z]+", "_", str(raw_value).casefold()).strip("_")
    aliases = {
        "transmission": ProjectType.TRANSMISSION_LINE,
        "transmission_line": ProjectType.TRANSMISSION_LINE,
        "line": ProjectType.TRANSMISSION_LINE,
        "substation": ProjectType.SUBSTATION,
        "reconductoring": ProjectType.RECONDUCTORING,
        "transformer": ProjectType.TRANSFORMER,
        "distribution": ProjectType.DISTRIBUTION,
        "generation": ProjectType.GENERATION,
    }
    project_type = aliases.get(normalized, ProjectType.UNKNOWN)
    state = (
        VerificationState.VERIFIED_RULE
        if project_type is not ProjectType.UNKNOWN
        else VerificationState.UNRESOLVED
    )
    evidence_id = bind("project_type", raw_value, project_type.value, state=state)
    return EvidenceValue[ProjectType](
        value=(project_type if project_type is not ProjectType.UNKNOWN else None),
        raw=str(raw_value),
        state=state,
        evidence_ids=[evidence_id],
    )


def _text_value(field: str, raw_value: object, bind) -> EvidenceValue[str] | None:
    text = _optional_text(raw_value)
    if text is None:
        return None
    evidence_id = bind(field, raw_value, text)
    return EvidenceValue[str](
        value=text,
        raw=str(raw_value),
        state=VerificationState.VERIFIED_RULE,
        evidence_ids=[evidence_id],
    )


def _column_lookup(headers: list[str]) -> dict[str, int]:
    lookup = {}
    for index, header in enumerate(headers):
        canonical = _canonical_header(header)
        if canonical:
            lookup[canonical] = index
    if "name" not in lookup:
        raise ValueError("starter workbook must contain a project-name column")
    return lookup


def _canonical_header(header: object) -> str:
    normalized = re.sub(r"\s+", " ", str(header).strip().casefold())
    for canonical, aliases in HEADER_ALIASES.items():
        if normalized in aliases:
            return canonical
    return normalized.replace(" ", "_")


def _optional_text(value: object) -> str | None:
    if value is None or not str(value).strip():
        return None
    return normalize_project_name(value)


def _evidence_id(source_version_id: str, sheet: str, row: int, field: str) -> str:
    seed = f"{source_version_id}|{sheet}|{row}|{field}"
    return f"EV-{hashlib.sha256(seed.encode()).hexdigest()[:14]}"


def _serializable_value(value: object) -> object:
    if hasattr(value, "value"):
        return value.value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value
